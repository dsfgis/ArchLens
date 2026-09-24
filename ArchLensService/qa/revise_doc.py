from pathlib import Path
from copy import deepcopy
import re, json, hashlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent.parent
doc=Document(ROOT/'qa/source_v1.0.docx')
body=doc._element.body
heads={int(m.group(1)):p for p in doc.paragraphs if p.style.name=='Heading 1' and (m:=re.match(r'^(\d+)  ',p.text))}
appendix=next(p for p in doc.paragraphs if p.text.startswith('附录 A'))
def p(text,style='normal'):
    x=doc.add_paragraph(text,style=style)
    x.paragraph_format.keep_with_next=False
    return x
def table(rows):
    n=len(rows[0]); t=doc.add_table(rows=0,cols=n);t.autofit=False
    widths=([1.6,5.5] if n==2 else [1.3,2.1,3.7] if n==3 else [1.0,1.6,2.2,2.3])
    if rows[0][0]=='契约':widths=[1.3,2.8,3.0]
    if rows[0][0]=='工具':widths=[2.0,2.0,3.1]
    if rows[0][0]=='方法和资源':widths=[2.6,1.95,2.55]
    if rows[0][0]=='表组':widths=[1.75,2.6,2.75]
    if rows[0][0]=='对象和状态':widths=[1.7,2.25,3.15]
    if rows[0][0]=='from 到 to':widths=[1.75,2.05,3.3]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        rr=t.add_row(); pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for c,txt,w in zip(rr.cells,row,widths):
            c.width=Inches(w)
            if rows[0][0] in ['契约','表组'] and i>0:txt=txt.replace('/', ' / ')
            c.text=txt;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in c.paragraphs:
                pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.keep_with_next=False
                for r in pp.runs:r.font.size=Pt(9);r.bold=i==0
    p('')
    return t
def render(md):
    lines=md.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('```'):
            buf=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):buf.append(lines[i]);i+=1
            x=p('\n'.join(buf));x.paragraph_format.keep_together=True
            for r in x.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
        elif line.startswith('## '):
            x=p(line[3:],'Heading 2');x.paragraph_format.keep_with_next=True
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([a.strip() for a in lines[i].strip().strip('|').split('|')]);i+=1
            table(rows);continue
        else:p(line)
        i+=1
def put(n,md,replace=False):
    start=heads[n]._p;end=heads[n+1]._p if n<20 else appendix._p
    if replace:
        el=start.getnext()
        while el is not end:
            nxt=el.getnext();body.remove(el);el=nxt
    before=set(body);render(md)
    new=[e for e in body if e not in before]
    for el in new:end.addprevious(el)

# Retain the original twenty-chapter organization and document furniture.
for t in doc.tables[:2]:
    for row in t.rows:
        for c in row.cells:
            if '2026年9月11日' in c.text:c.text=c.text.replace('2026年9月11日','2026年9月12日')
            if '1.0  立项与技术评审' in c.text:c.text='1.1  实施设计修订  基于原 v1.0'
            if '提议评审版' in c.text:c.text=c.text.replace('提议评审版','实施设计评审版')
replacements={
'若 SQL 改为 global_id AS event_id':'兼容判断必须区分现有事实和拟议方案。global_id AS event_id 只能证明查询输出标签可能保持；还需核验 Java 映射、JSON 契约、类型和值语义。拟议别名尚未实施时，changeRequired 保持 UNKNOWN，并记录条件性结论；具体状态和合并规则见第9章。',
'任务持有租约并定期续期':'任务通过 PostgreSQL 队列领取，持有租约并定期续期。重领时递增 attemptEpoch，旧 epoch 的续期与写入一律拒绝。幂等批次、取消和封存协议见第13章；分析器语法错误形成诊断，不重复重试。',
'通过 information_schema 与 pg_catalog':'通过 information_schema 与 pg_catalog 采集 schema、表、列、类型、默认值、约束、索引、视图定义和登记依赖。[R5] 元数据采集使用同一连接上的 REPEATABLE READ READ ONLY 事务；默认 READ COMMITTED 的多条查询不能保证同一视图。[R9] 固定 search_path 和语句超时30秒，整批采集上限60秒。对象消失或并发 DDL 导致查询失败时回滚整批，最多重采2次，禁止拼接不同尝试结果。SQL只允许内置模板；只读事务不代替查询白名单。Git 与数据库不是原子快照，须以发布清单关联，未核实时 alignment=UNVERIFIED。',
'模型引用不存在的 evidenceId':'模型引用不存在的 evidenceId、不同项目证据或不同快照证据，由服务端验证器强制拒绝；证据与结论的语义相关性由 Reviewer 辅助检查，并保留人工复核。模型不通过不是事实被自动推翻，模型通过也不代表证明成立。',
'状态区分 DRAFT':'状态以第11章的 Case 与 InvestigationRun 双层定义为准。',
}
for pp in doc.paragraphs:
    for prefix,value in replacements.items():
        if pp.text.startswith(prefix):pp.text=value
    if pp.text.startswith('读证据时检查项目权限'):
        pp.text='读证据时检查项目权限、snapshotId、文件哈希、行列范围和内容摘要，默认显示固定快照。不存在、越权和跨快照引用由服务端强制拒绝；证据是否支持语义结论由 Reviewer 辅助评估，不能保证模型零漏判。与原分支最新文件比较必须显式请求。'
    if pp.text.startswith('调查固定 snapshotId'):
        pp.text='调查固定 snapshotId、overlayRevision、ruleVersion、promptVersion、modelConfigVersion。数据源扫描不自动切换调查快照。API 返回 traceId；Worker 记录 scanId 和 stage；监测解析成功率、未知符号率、投影延迟、截断率及模型成本。'
    if pp.text.startswith('QUEUED -> COLLECTING'):
        pp.text='Scan: QUEUED -> COLLECTING -> ANALYZING -> VALIDATING\nSnapshot: BUILDING -> SEALED -> PROJECTING -> PUBLISHED\nScan terminal: SUCCEEDED / PARTIAL / FAILED / CANCELLED\nSnapshot quality: COMPLETE / PARTIAL'
    if pp.text.startswith('扫描完成后先写 manifest'):
        pp.text='扫描先封存包含各仓库提交、schemaHash、配置与依赖摘要、解析器版本、诊断和记录摘要的 manifest。投影校验通过才发布 PUBLISHED；quality 单独为 COMPLETE 或 PARTIAL。关键来源整体不可用时扫描失败且不发布，允许发布的部分能力必须逐项列明。发布协议见第13章。'
    if pp.text.startswith('User / Reviewer'):
        pp.text='User / Reviewer\n       | HTTPS and project authorization\n       v\nVue Web ---> Spring Boot API ---> Case Service\n                  |                    |\n          Scan Coordinator       Investigator <--> Model Gateway <--> LLM\n                  |                    |\n          Java / Node Workers     Tool Gateway\n          PG Schema Collector          |\n                  |              Graph / Evidence / Blast\n          Normalize and Validate       |\n                  |                    |\n                  v                    v\n           PostgreSQL facts and cases and outbox\n                  | immutable projection\n                  v\n              Neo4j graph'
        for r in pp.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
    if pp.text.startswith('Scan: QUEUED'):
        for r in pp.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
    if pp.text.startswith('保存源码全文不是必要前提'):
        pp.text='证据片段以脱敏内容及 snippetHash 保存，原始 sourceHash 记录字节身份。完整源码位于隔离扫描卷，在被有效 Case 引用期间保留；原始字节可用时复验原始哈希，否则仅校验片段哈希并标识 originalAvailable=false。被合规删除的证据标 unavailable，保留删除审计，不继续显示完整可复查。'

put(1,r'''
## 实施修订范围
本次修订保留原有总体架构，补齐对象级调查加有限字段映射、变更假设、图契约、传播与评分规则、模型调用、任务恢复、接口闭环和验收工作包。文档内版本为1.1，交付文件沿用原文件名以便替换。所有性能和准确率均是待实施验证的门槛，不代表已通过测试。
REQ01 固定多来源快照并生成可定位证据；REQ02 在声明的解析能力内发现依赖并保留未知；REQ03 对五类变更进行有界调查和条件性修改判断；REQ04 使用受限模型工具生成有引用的建议；REQ05 支持可恢复调查、修订评审和导出；REQ06 执行权限、隔离、保留和审计；REQ07 通过独立黄金集与恢复验证。第16章工作包和第20章验收项引用这些需求编号。
''')
put(3,r'''
## 首版实施决策
ADR 05：首版使用对象级依赖图加有限字段映射；支持直接结果映射、简单属性赋值、显式 JSON 命名和前端直接属性访问。反射式复制、复杂表达式、自定义序列化与不完整上下文不能自动得出 NO。此边界同时约束 UI 文案、规则与黄金集。
ADR 06：单机开发部署使用 API、Java分析Worker、Node分析Worker、PostgreSQL、Neo4j；任务队列复用 PostgreSQL，不额外引入消息队列。HTTP 请求与扫描进程分离。企业模型网关是外部依赖，断开时仍可创建结构化调查并输出图报告。
ADR 07：Case 审批与执行状态分离；已发布快照、终结的 CaseRevision 和 Run 结果不可改写。人工 overlay 不修改基础事实，只对固定修订的调查视图生效。
''')

put(6,r'''
本章定义跨 Java、Node、数据库和模型工具的契约版本 archlens-contract-v1，覆盖 REQ01、REQ02、REQ03。线上 JSON 使用 camelCase，数据库使用 snake_case；同一字段只允许一种类型。以下 R 表示必填，O 表示可选；未列出的字段拒绝，只有 attributes 对应的类型专用 schema 允许扩展。
## 身份和公共数据类型
业务ID使用服务端分配的UUID；事实ID为64位小写十六进制SHA256。H(parts) 表示对无歧义JSON数组的UTF-8字节哈希：对象键按代码点排序，数组保持规定顺序，无空白；标识符按所属语言规范处理，不对路径统一转小写。schemaVersion是字符串，时间使用UTC RFC3339，分数decimal取0到1，计数为非负整数。位置行列从1开始，end不包含，byteOffset为原始字节偏移。
nodeId = H([projectId,snapshotId,type,logicalKey])。logicalKey以sourceId开头；Java加入module及全限定类和参数签名，数据库加入databaseKey/schema/table/column，前端加入仓库路径和符号键。同名双仓文件及不同数据库同名表不会合并。改名产生新logicalKey；跨快照匹配存独立identity_mapping，不作为同快照传播边。
relationId = H([projectId,snapshotId,fromId,toId,kind,conditionHash,bindingKey])；绑定键区分相同端点的不同字段映射。重复采集同一关系合并evidenceIds并排序；不同分支保留独立边。evidenceId = H([projectId,snapshotId,sourceId,sourceHash,location,kind,producerVersion])。pathId = H([snapshotId,overlayRevision,targetId,orderedRelationIds,semanticState])。
| 契约 | 字段与类型 | 校验及语义 |
| SourceManifest | R projectId/snapshotId UUID；R sources[]；R analyzerVersions/configHash/dependencyHash；R quality/diagnostics/counts/contentDigest | sources每项R sourceId、kind、contentHash；REPO需commitSha，DATABASE需databaseKey/schemaHash/collectedAt；记录alignment和capabilities |
| NodeIR | R schemaVersion/projectId/snapshotId/nodeId/logicalKey/type/name/sourceId/sourceRef/attributes | sourceRef含sourceHash及location；库对象以metadataKey定位；type校验专用attributes；scope完全一致 |
| EdgeIR | R relationId/fromId/toId/kind/condition/bindingKey/evidenceIds/certainty/confidence/producerVersion；含公共scope | 端点必须存在；CONFIRMED至少1条有效证据；confidence不得由模型填写；空条件为规范化true |
| EvidenceIR | R evidenceId/sourceId/sourceHash/kind/producerVersion/location/validation/certainty/condition；O snippet/snippetHash/redaction/derivedFrom | 同scope；源码按行列和字节定位，数据库按对象键和采集查询定位；删除状态不删除审计 |
| DiagnosticIR | R diagnosticId/sourceId/code/message/location；O candidateNodeIds/capability | 未解析符号、超限、权限缺失均在此记录，不能把不存在端点写为确定边 |
| AnalysisBatch | R scanId/attemptEpoch/batchKey/schemaVersion/nodes/edges/evidence/diagnostics/contentHash | 同键同内容可重放，异内容409；原始批次可有未绑定引用，封存前完成归并 |
| ImpactResult | R targetId/nodeId/dependencyPresent/changeRequired/certainty/impactScore/risk/pathIds/evidenceIds/unknownIds/conditions | changeRequired只允许YES/NO/UNKNOWN；conditionedOutcome可为NO；NO需规则完备证据，截断不自动降低风险 |
## 节点属性与关系字典
节点枚举为SYSTEM、MODULE、CLASS、METHOD、FIELD、TABLE、COLUMN、VIEW、SQL_STATEMENT、API、VUE_PAGE、COMPONENT、FUNCTION、CONFIG、JSON_FIELD、FRONTEND_ACCESS。JOB、TOPIC、GIS_LAYER、EXTERNAL_API为禁用的扩展类型，启用需契约升级。统一用VUE_PAGE，不再使用PAGE。
METHOD.attributes含signature、declaringType、parameterTypes、returnType和module；COLUMN含dbType、nullable、default、isUnique及约束引用；API含serviceKey、HTTP方法、pathTemplate、params、headers、consumes、produces、requestFields和responseFields。[R2] JSON_FIELD含apiId、REQUEST或RESPONSE、JSON Pointer、类型、required及sourceFieldId；FRONTEND_ACCESS含functionId或componentId、propertyPath和accessMode。未知类型允许null并带diagnosticId，不能猜测为string。
| from 到 to | kind | 事实含义与传播方向 |
| METHOD到METHOD；VUE_PAGE或FUNCTION到FUNCTION | CALLS | 调用者依赖被调用者，影响反向；接口声明与唯一可解析实现用实现绑定证据连接，歧义仅候选 |
| METHOD到COLUMN | READS 或 WRITES | SQL绑定的方法访问列，反向；保留statementId及分支条件 |
| API到METHOD；FUNCTION到API | HANDLED_BY；CONSUMES | 端点处理、请求消费，反向 |
| FIELD到COLUMN | MAPS_TO | 查询结果映射到属性的投影关系；保留结果标签及resultMapId，反向 |
| FIELD到FIELD；JSON_FIELD到FIELD | DERIVED_FROM；SERIALIZES_FROM | 直接赋值或已识别序列化映射，反向；复杂转换标候选 |
| FRONTEND_ACCESS到JSON_FIELD | READS_FIELD | 精确属性读取，反向；再由读取点归属聚合到页面，不遍历全仓CONTAINS |
| TABLE到COLUMN；COLUMN到COLUMN | CONTAINS；FK_TO | 归属不传播；外键的被引用列变化时反向 |
查询结果标签、Java属性和JSON名称分开保存；请求字段使用独立的DESERIALIZES_TO关系（FIELD到JSON_FIELD，含请求证据），首版只用于直接参数绑定。返回整个DTO而没有可识别字段映射时，保留对象级影响并输出FIELD_LINEAGE_GAP。相同API路径必须结合服务与全部已支持条件匹配，不同produces不合并。
## 候选关系与人工修正
certainty=CONFIRMED或INFERRED；UNKNOWN用于Diagnostic和影响结论，不是伪造端点。confidence初始规则值：精确声明或显式绑定1.0，受控唯一候选但条件未决0.6；不能按名称相似升级为CONFIRMED。未落定端点写unresolved_reference；候选端点存在时可写INFERRED边。两类均计入coverage缺口。
overlay含overlayId、projectId、snapshotId、revision、operation、targetRef、reason、evidenceRefs、actorId和createdAt；operation允许ADD_CANDIDATE、REJECT_EDGE、ANNOTATE。人工新增边默认INFERRED，不能直接升级解析器CONFIRMED；被否定边保留基础事实并在有效调查视图中排除，报告说明修正。Case固定overlayRevision；新增overlay需要新调查，不能重算旧审批结果。
## 合法性与落库约束
例如NodeIR(type=COLUMN, sourceId=dbA, logicalKey=dbA/public/device/event_id)与EdgeIR(kind=READS, fromId=存在的METHOD, toId=该列, certainty=CONFIRMED, evidenceIds=[有效SQL证据])是合法记录。缺少证据的确定边返回EVIDENCE_REQUIRED；指向另一快照返回SCOPE_MISMATCH；type=PAGE返回UNSUPPORTED_NODE_TYPE；同ID异内容返回IDENTITY_CONFLICT。以上四种非法样例必须进入契约测试，不能被自动修复成有效事实。
''',True)

put(7,r'''
## 解析能力矩阵
以下矩阵是首版可验收范围；未覆盖模式保留对象依赖或候选，并必须生成诊断，不让模型补写确定边。每项在黄金样例中记录输入文件、位置、期望节点和边。
| 能力ID | 首版支持 | 明确降级条件 |
| PAR01 Java符号 | 直接调用、重载签名、继承声明、唯一可解析注入目标 | 缺依赖、多个Bean、反射、动态代理目标不唯一为UNRESOLVED_SYMBOL或AMBIGUOUS_BINDING |
| PAR02 SQL映射 | 显式resultMap、读取配置后的驼峰映射、静态别名、常规注解SQL | 动态标识符、自定义TypeHandler语义、复杂映射为DYNAMIC_IDENTIFIER或MAPPING_GAP |
| PAR03 动态SQL | include展开及循环检测，if/choose/foreach/trim/where/set；最多64个规范化分支 | 超预算保留已有引用加BRANCH_LIMIT；不声称未解析分支无依赖 |
| PAR04 API契约 | 常规MVC映射、一层组合注解、直接返回DTO、显式JsonProperty和JsonIgnore | 全局未知命名策略、自定义序列化、反射复制、泛型结构未绑定为SERIALIZATION_GAP |
| PAR05 有限字段映射 | 直接赋值、直接getter返回、静态结果标签、DTO属性映射 | BeanUtils、MapStruct生成代码不可用、算式和分支值转换为FIELD_LINEAGE_GAP |
| PAR06 Vue请求 | axios/fetch及最多2层本地封装，静态baseURL，字面量或单参数路径 | 拦截器重写、动态host、复杂拼接为DYNAMIC_URL；路由字面量import可识别，计算路径不支持 |
| PAR07 前端字段 | response.data.eventId、已绑定解构、模板直接属性访问 | 动态下标、spread、对象整体转交未知组件为FRONTEND_FIELD_GAP，禁止由未发现读取推出NO |
classpath由管理员上传的依赖清单和允许的制品缓存提供，记录坐标、绝对解析后的缓存引用与SHA256；不运行来源仓库Maven插件或注解处理器。JDK模块由选定运行时解析。Lombok源声明可建立字段，但未实际存在的生成方法不假装已解析；可接收已授权且与提交绑定的生成源码或字节码，否则降级。依赖缺失不能阻止本文件语法树提取，但相关符号边不得确认。
SQL规范化先保留条件树，再形成可解析的完整语句骨架；值占位符转解析器接受的参数节点，foreach使用一个符号元素表示集合语法。不能把孤立AND片段直接当完整SELECT。SourceMap保存展开片段到原文件的多段位置，include证据同时指向使用处和定义处。[R3][R10] 不执行OGNL。多个Mapper数据源绑定必须来自受控配置；不唯一时不将SQL绑定到任意库。
source配置必须显式包含moduleRoots、encoding、activeProfiles、mapperLocations、dataSourceBindings、serviceRouteMappings、frontendAliases及requestWrappers。首版每来源固定一个分析环境；多个Profile分别建快照。route映射明确服务、contextPath、gatewayPrefix、baseURL与代理重写，缺一项若影响唯一匹配则保留候选。
## 字段级判定的完备性
fieldCoverage对每个API和消费页面标记COMPLETE_FOR_SUPPORTED_PATTERN或PARTIAL，并列出所有逃逸点。只有整个受影响字段映射路径在支持模式内、类型和语义兼容且无未知出口，才允许规则产生NO。对象级链路始终展示；未使用某字段的判断仅在页面所有消费点已覆盖时有效，不推广到其他页面或生产运行时。
''')

put(9,r'''
Blast Radius计算待评估对象，修改结论由变更语义与证据另行确定。规则版本统一为ruleVersion，包含传播、兼容、置信度及风险评分子规则；不再另设含义重叠的policyVersion。所有结果满足REQ03。
## 变更输入与条件性结论
ChangeSpec必填targetId、kind、before、after、stage、compatibility、assumptions；before必须与快照目标核对，冲突返回BEFORE_MISMATCH。kind为COLUMN_RENAME、COLUMN_DROP、COLUMN_TYPE_CHANGE、API_CONTRACT_CHANGE或METHOD_SIGNATURE_CHANGE；仅接受一个目标，批量改动分为多个关联Case。
before和after采用kind专用结构：列含name/type/nullable/identityMeaning，删除时after=null；API含HTTP方法、路径和请求响应字段变化列表；方法含完整signature及参数变化。stage为PROPOSED或APPLIED。compatibility每项含kind、scopeIds、state、evidenceIds、conditions；state为PROPOSED、VERIFIED或REJECTED。VERIFIED只能由规则核对实际固定快照或关联验证快照及验证记录后赋值，模型与用户布尔值不能直接设定。
验证另一个快照时，compatibility使用可选validationRef，包含validationSnapshotId、recordId、mappingRefs、checkedConditions和contentHash。该记录属于比较证明，不加入基准快照的Evidence Chain；普通evidenceIds仍只能引用当前快照。跨版本节点映射必须显式确认，不能按同名合并。stage=APPLIED时before与所关联基准快照核对，after与验证快照核对；两者缺一则拒绝声称已验证实施。
```json
{
  "targetId": "<columnNodeId>", "kind": "COLUMN_RENAME",
  "before": {"name":"event_id","type":"text",
             "nullable":false,"identityMeaning":"device-key"},
  "after": {"name":"global_id","type":"text",
            "nullable":false,"identityMeaning":"device-key"},
  "stage":"PROPOSED",
  "compatibility":[{"kind":"SQL_ALIAS","scopeIds":["<mapperId>"],
    "state":"PROPOSED","evidenceIds":[],
    "conditions":["retain JSON eventId and value semantics"]}],
  "assumptions":[]
}
```
dependencyPresent表示在本调查视图发现依赖；changeRequired为YES、NO或UNKNOWN。YES表示在给定变更生效条件下存在已证实的需改消费点；NO只对本快照及支持范围成立。拟议兼容方案只能得到UNKNOWN加conditionedOutcome=NO及条件，不伪造当前事实。纯物理改名与业务标识替换必须分开；名称相同不证明值可互换。
## 传播状态与规则表
状态为nodeId、depth、strength、semanticKind、fieldRef、certainty、boundaryState和orderedPath。semanticKind为PHYSICAL_SCHEMA、VALUE_TYPE、FIELD_CONTRACT、METHOD_SIGNATURE、OBJECT_DEPENDENCY。边权只排序调查优先级，不决定YES。保留完整依赖调查通道和修改判定通道：已验证兼容边界仅停止对应语义的修改传播，不删除下游依赖。
| 规则 | 输入与边 | 判定及后续语义 |
| BR01 列改名 | PHYSICAL_SCHEMA经READS/WRITES反向 | 静态SQL仍引用旧列为YES；仅对应SQL兼容证据完整才停止修改传播；经CALLS上溯默认OBJECT_DEPENDENCY和UNKNOWN |
| BR02 列删除 | PHYSICAL_SCHEMA经READS/WRITES/MAPS_TO/FK_TO反向 | 已确认引用为YES；存在替代字段方案不等于兼容已证实；转相关FIELD_CONTRACT |
| BR03 类型变化 | VALUE_TYPE经上述列关系及字段映射反向 | 相同类型或白名单无损转换可局部兼容；收窄、nullable收紧存在明确冲突为YES；无类型证据为UNKNOWN |
| BR04 API契约 | FIELD_CONTRACT经CONSUMES/READS_FIELD反向 | 删除或改名被直接读取字段为YES；完整消费范围证明未读取该字段才为NO；请求和响应分别匹配 |
| BR05 方法签名 | METHOD_SIGNATURE经CALLS反向 | 已解析调用与新参数数量或类型不兼容为YES；重载兼容需保留旧签名证据；上层对象默认UNKNOWN |
| BR06 字段映射 | 经DERIVED_FROM/SERIALIZES_FROM/MAPS_TO反向 | 保留fieldRef并验证名称、类型、值语义；复杂转换或未知出口降级UNKNOWN；DESERIALIZES_TO按请求流单独处理 |
| BR07 结构聚合 | CONTAINS或读取点到页面归属 | 普通CONTAINS不扩散；FRONTEND_ACCESS只聚合已记录owner页面，结果来自其消费点集合 |
数据库类型兼容首版采用明确白名单：同类型同约束为等价；smallint到integer、integer到bigint仅在Java接收类型和JSON表现也兼容时允许通过；varchar(n)扩大仅证明数据库容量不收紧。其他类型变化，包括text到uuid和时区语义变化，默认UNKNOWN。METHOD_SIGNATURE仅对已绑定实参和Java允许的转换判断，不自行重现完整编译器。
## 多路径合并与有界算法
每个消费点独立判定：任何CONFIRMED路径给出YES则节点为YES；没有YES但有UNKNOWN、未验证条件、候选冲突或相关覆盖缺口则为UNKNOWN；只有所有相关消费点均已验证兼容或无须改动且完整展开才为NO。截断后已证实YES保留，其余不得因已保存路径全兼容而升级NO。一个接口兼容不会覆盖另一个接口的不兼容结果。
```text
queue = [State(target, 0, 100, initialSemantic, [])]
while queue and workUnits < maxWorkUnits and beforeDeadline():
    s = popBy(strength desc, logicalKey, relationId, stateHash)
    for edge in sortedAllowedIncoming(s):
        if edge.scope != fixedScope: fail(SCOPE_MISMATCH)
        if nextNode(edge) in s.pathNodes: continue
        r = applyRule(s, edge, changeSpec, fixedCompatibility)
        mergeConsumerJudgment(r)  # independent of display paths
        saveBestExplanationPaths(r, limit=3)
        if canExpand(r): enqueueUnlessExactDuplicate(r)
        else: recordBoundaryReason(r)
return aggregateJudgments(), coverage, truncationReasons
```
首版不做跨语义优势剪枝，只删除路径和完整状态完全相同的重复项，避免不同visited集合被错误合并。maxDepth=8、maxNodes=5000、maxWorkUnits=50000，在线墙钟上限5秒；深度边界只在存在可继续遍历的边时标截断。循环边单独统计，不据此声明未知。每节点3条仅是展示上限；判断聚合独立处理所有已访问路径。允许候选传播，但不得将其升级确定结论。固定工作预算的重算与墙钟超时回放见第20章。
## 调查强度与证据质量
边权READS/WRITES/MAPS_TO/FK_TO及有限字段映射为1.0，CALLS/HANDLED_BY/CONSUMES为0.9；归属聚合不增加跳数。I(target)=100；h≥1时I(path)=100×连乘边权×0.85^(h-1)，I(node)取路径最大值。C(path)取边confidence最小值，C仅描述规则证据质量，不是概率且不乘入风险。计算使用未舍入数，显示I保留2位小数，风险整数采用HALF_UP。
原对象链示例列到Mapper一跳I=100，列到页面六跳I=100×0.9^5×0.85^5约为26.200355，显示26.20。新增字段映射路径按实际边数计算，不继续套用对象链六跳。边界和UNKNOWN不会因距离较远变为低风险。
## 风险输入和评分规则
RiskInput包含B/K/D的value或[min,max]、sourceKind、sourceRef、ruleId和confirmedBy。sourceKind为RULE、PROJECT_CONFIG或HUMAN_CONFIRMED；模型仅生成proposal，不得写入正式数值。缺失B或D为[0,1]；缺失K也为[0,1]，不能默认低风险。
| 因子 | 首版赋值 | 必要依据 |
| B破坏程度 | 删除被消费字段或已确认不兼容签名1.0；直接SQL列改名0.8；已验证兼容0；仅候选则区间 | BR规则ID、消费点证据；同一节点取最高已证实程度，未知可能性扩展上界 |
| K业务关键性 | 核心生产或关键对外服务1.0；重要内部流程0.7；一般辅助0.3 | 项目或资产配置，由责任人确认；无配置则区间 |
| D数据与恢复 | 已证实不可逆1.0；需映射回填或恢复演练0.8；有验证可逆步骤0.3；无数据变更且已确认0 | 已关联的验证或人工确认记录；仅凭源码不能声称历史数据安全 |
R=HALF_UP(100×(0.40B+0.35K+0.25D))；有区间时分别计算单调的上下界并显示等级区间，不仅展示下界。R≥80为CRITICAL，60到79为HIGH，35到59为MEDIUM，其余LOW。已证实删除关键唯一标识且缺少映射恢复方案，风险下限提高到60并记录overrideRuleId；只是疑似关键标识时提出待确认，不伪造触发条件。示例B=1、K=0.8、D=0.9来自已确认配置时R=91；置信度0.6不降低R。
''',True)

put(10,r'''
Agent在确定性工具之上理解需求、补证、归纳风险和生成计划，满足REQ04。原始源码、证据和模型输出都是数据，服务器掌握状态、权限、预算及落库。一个模型可承担不同阶段角色，不要求独立模型进程。
## 模型网关和结构化输出
ModelGateway.generate(ModelRequest)返回ModelResponse。请求必填runId、stage、modelConfigRef、promptVersion、messages、allowedToolSchemas、responseSchema、deadline和maxOutputTokens；消息只含最小授权证据包及稳定引用。响应为TOOL_CALLS、STRUCTURED_RESULT、REFUSAL或ERROR，含providerRequestId、usage、finishReason及payload。厂商字段由适配器转换，不渗透到核心Case模型。
配置版本含provider、modelId、approvedEndpointRef、credentialRef、contextLimit、input/output token上限、价格配置版本和超时；密钥本身不进入Case。价格不可用时费用显示UNKNOWN，token与次数仍限制。输入含证据目录、变更、已验证图结果、未知项和剩余预算，不能把整个仓库或未授权片段塞入上下文。
| 阶段 | 输入与模型职责 | 输出及可用工具 |
| INTAKE | 自然语言及候选目标，识别变更语义与缺失条件 | IntentResult含targetQuery、proposedChange、questions；仅search_architecture |
| INVESTIGATE | 已确认ChangeSpec、固定快照、coverage；选择下一调查步骤 | 工具调用或FindingsDraft；只读工具；每条finding含claim/evidenceIds/pathIds/certainty/assumptions |
| PLAN | 已验证影响、规则风险、未知项；组织变更和测试建议 | PlanDraft含risks解释、planTasks、testScopes；不得改分数，可读取证据及save_case_draft |
| REVIEW | 草案、引用目录、反例、验证器结果 | ReviewResult含PASS/NEEDS_EVIDENCE/NEEDS_HUMAN、issues、findingIds、reason；只读工具 |
提示词四个模板分别版本化。共同约束为仅引用输入或工具返回ID、把推测标为假设、不得变更快照/租户/评分、遇到未知不得输出无影响、输出必须符合schema。每条PlanTask必填id/title/reasonRefs/dependsOn/ownerRole/preconditions/deliverable/validation/rollback；TestScope必填impactedNodeIds/riskIds/testLevel/scenario/expectedResult/existingTestRefs/newTestRequired。dependsOn必须组成无环图，测试索引为空时existingTestRefs=[]且newTestRequired=true。
## 工具契约
公共ToolContext由服务端注入projectId、snapshotId、overlayRevision、actorId、runId和traceId；调用参数禁止重复提供这些字段。所有工具结果统一为data、diagnostics、truncated、nextCursor；失败为error.code/message/retryable/details。limit默认50、最大200；字符串查询最长256，evidence单次总长度上限12000字符，路径深度最大8，批量ID最多50。
| 工具 | 必填及可选参数 | data结构 |
| search_architecture | R query；O types/limit/cursor | items含nodeId/type/name/logicalKey/matchReason；歧义不自动选择 |
| get_neighbors | R nodeId/direction/kinds；O limit/cursor | nodes、edges、boundaryReasons；direction为IN/OUT/BOTH |
| find_dependency_path | R fromId/toId；O maxDepth/maxPaths | paths含pathId/nodeIds/relationIds/evidenceIds；maxPaths默认3最大10 |
| get_source_evidence | R evidenceIds；O maxChars | evidence含id/location/snippet/validation/redaction/originalAvailable |
| query_database_schema | R objectIds | objects含对象键、列类型及约束；只读已采集元数据，禁止SQL文本 |
| calculate_blast_radius | R changeSpec | ImpactResult[]、paths、riskInputs、unknowns、coverage、ruleVersion；规则版本由Run固定 |
| search_code | R literal；O allowedPaths/limit/cursor | matches含sourceId/path/location/sourceHash；仅候选证据，路径与白名单取交集 |
| save_case_draft | R draftVersion/findings/planTasks/testScopes | draftVersion、validationIssues；仅写Run草案，不能改图、分数、正式修订或审批 |
统一参数入口采用JSON Schema校验且additionalProperties=false。scope不存在或不可见返回NOT_FOUND；无工具执行权返回FORBIDDEN；参数错误INVALID_ARGUMENT；快照未发布SNAPSHOT_NOT_READY；图不可用GRAPH_UNAVAILABLE。模型调用calculate_blast_radius只传changeSpec，不再使用旧targetId/change/policyVersion形状。服务端将第9章完整ChangeSpec和Run规则版本组装到引擎请求。
## 调用循环和预算
```text
validateIntentOrWaitForClarification()
persistRunContextAndBudgets()
while stageNotFinished and budgetAvailable:
    response = modelGateway.generate(authorizedContext)
    validateResponseSchema(response)
    if TOOL_CALLS:
        for call in response.callsInOrder:
            validateStageAndArguments(call)
            result = executeWithCallIdempotency(call)
            persistToolResultAndBudget(result)
        continue
    if STRUCTURED_RESULT:
        validateReferencesAndDomainRules(response)
        checkpointAndAdvanceOrRequestEvidence()
    else: terminateOrFallbackWithDiagnostics()
finalizeImmutableRevisionThroughCaseService()
```
首版串行执行工具。每Run最多12次模型请求、30次工具执行、2轮补证，累计活跃时间120秒；用户等待澄清不消耗活跃时间，预算不因恢复重置。单模型请求上限30秒并受剩余deadline约束，默认输入预算32000 tokens、累计输出预算12000 tokens，实际不超过provider上下文；预算值是首版可调策略，不是性能保证。
同一response内重复toolCallId仅返回缓存结果；同参数读调用可复用run内缓存，同工具同参数连续3轮且无新证据时终止为NO_PROGRESS。save_case_draft幂等键为runId/toolCallId；同键异内容拒绝。模型网络错误最多重试1次并计入请求预算；未知是否计费时记录usageUncertain。鉴权、拒答不盲重试；JSON或参数错误允许1次带校验信息的修复请求，仍失败结束为PARTIAL或FAILED。
## 复核与降级
程序校验引用存在性、scope、哈希状态、PlanTask依赖和字段类型；这类错误未修复的结论不得进入正式报告。Reviewer检查语义支持和反例，不能保证判断完美；争议结论进入NEEDS_HUMAN并保持UNKNOWN。已确认事实仅经人工overlay改变调查视图。补证耗尽保留合格结果，缺失项进入unknowns，resultQuality=PARTIAL。
模型不可用且有已确认ChangeSpec时，CaseService仍可调用图算法生成纯图报告；没有确认目标或变更语义则等待澄清，不能猜测后降级。图不可用返回失败或历史报告回放，不允许模型补图。保存工具输入输出摘要、模型实际响应与版本引用以便审计；受限片段按权限与保留策略存储。
''',True)

put(11,r'''
Investigation Case是调查和评审聚合，InvestigationRun是一次执行尝试序列。CaseRevision保存固定输入与最终结果；RunDraft保存执行中的可变草案。业务审批不代表代码已修改或上线，满足REQ05。
## 聚合与版本
CaseRecord含caseId/projectId/title/ownerId/latestRevision/recordVersion/reviewStatus；CaseRevision含caseId/revision/snapshotId/overlayRevision/changeSpec/assumptions/impacts/paths/risks/planTasks/testScopes/unknowns/coverage/ruleVersion/promptVersion/modelConfigVersion/resultQuality/runId/createdAt。版本引用均不可改写。创建Case先生成revision=1的输入修订；Run终结通过事务追加结果修订，不修改输入修订。
InvestigationRun含runId/caseId/baseRevision/recordVersion/state/stage/attemptEpoch/checkpoint/budgetUsed/draftVersion/error/startedAt/completedAt。RunDraft内容与工具日志可递增，但不是正式CaseRevision。每Case最多一个活跃Run。CaseRecord与Run各自的recordVersion是其API乐观锁ETag来源，与CaseRevision.revision、RunDraft.draftVersion分别命名。
## 状态迁移
| 对象和状态 | 触发及目标 | 约束 |
| Run QUEUED | worker领取到RUNNING | 保存epoch及起始上下文 |
| Run RUNNING | 歧义到WAITING_INPUT；完成到SUCCEEDED；可用但不完整到PARTIAL | WAITING_INPUT对应前端NEEDS_CLARIFICATION；每阶段先持久化checkpoint |
| Run WAITING_INPUT | 澄清到QUEUED；取消到CANCELLED | 新答案记录审计；影响输入时创建新输入修订并绑定Run，不改历史版本 |
| Run RUNNING或QUEUED | 取消到CANCELLING再CANCELLED；不可恢复错误到FAILED | 取消提交后禁止新增有效草案写入；保留已有检查点和诊断 |
| Case DRAFT | 完整或部分结果送审到READY_FOR_REVIEW | PARTIAL必须含unknowns和限制确认；纯失败不能送审 |
| Case READY_FOR_REVIEW | Reviewer批准到APPROVED或否决到REJECTED | 审批绑定caseId/revision/snapshotId，不作用于正在变化的草案 |
| Case APPROVED或REJECTED | 新调查结果生成后回到DRAFT；归档到ARCHIVED | 旧审批事件仍保留；ARCHIVED只读，解除归档需Admin审计操作 |
Run的SUCCEEDED/PARTIAL/FAILED/CANCELLED为终态；重新调查创建新Run，不复用已终结Run。REVIEW阶段模型PASS只表示生成流程完成，不能将Case设APPROVED。故障后从已持久化阶段恢复，网络返回但未落库的模型请求可能重新调用，记录重试和费用不确定性。
澄清只能补充当前目标和同一快照的语义；若改变targetId或snapshotId，取消当前Run并创建新调查。语义补充形成新输入修订时，Run的原始baseRevision保留，effectiveInputRevision指向新输入；清除依赖旧语义的草案计算并从相应阶段恢复，已消费预算不退还。只修改说明文字不重置图计算。
## 界面操作闭环
保留Architecture Graph、Blast Radius和Investigation Case三工作区。图按模块聚合并提供列表；影响列表显示YES/NO/UNKNOWN、条件性结论、风险区间和截断边界。字段级覆盖不完整必须可见。证据抽屉显示原始位置、固定版本、脱敏与权限状态。
Case页提供目标消歧、提交澄清、取消运行、查看历史修订、比较调查结果、人工标注、送审和导出。模型解释与算法事实分别标识来源。Viewer只查看已批准且可见的修订；Analyst可看项目草案。审批PARTIAL时Reviewer必须填写接受的限制及理由，不消除未知事实。
''',True)

put(12,r'''
API前缀/api/v1，JSON UTF-8。长任务返回202及Location；所有列表返回items/nextCursor。服务端以请求身份确定项目权限，ID命中后再次校验scope。下列R为必填，O为可选；payload引用第6、9、10、11章定义，不重复发明字段。
## 资源接口
| 方法和资源 | 请求 | 响应及关键行为 |
| POST /projects | R name；O description | 201 projectId；仅全局Admin可创建 |
| GET /projects；POST /projects/{id}/sources | 列表；R kind/config/credentialRef | 授权项目列表；201 sourceId，禁止原文密钥 |
| GET /projects/{id}/sources | O cursor/limit | 来源配置脱敏列表 |
| POST /projects/{id}/scans | R sourceIds/analyzerProfile/configRevision | 202 scanId/statusUrl；来源必须齐备 |
| GET /scans/{id}；POST /scans/{id}/cancel | 查询；取消reason | 状态阶段和诊断；202取消请求 |
| GET /projects/{id}/snapshots | O quality/cursor/limit | 已发布快照及manifest摘要 |
| GET /snapshots/{id}/nodes | O query/type/cursor/limit | 节点候选；空query按logicalKey排序 |
| GET /nodes/{id}/neighbors | R snapshotId；O direction/kinds/cursor | 第10章邻接结果 |
| POST /snapshots/{id}/paths | R fromId/toId；O maxDepth/maxPaths | 路径与证据；仅有界查询 |
| GET /evidence/{id} | 无 | 证据级授权；不可见时不返回片段 |
| POST /cases | R projectId/snapshotId/title；changeSpec或intentText二选一 | 201 caseId/recordVersion/revision；未消歧可保存DRAFT |
| POST /cases/{id}/investigations | R baseRevision | 202 runId；已有活跃Run返回409 |
| GET /runs/{id}；POST /runs/{id}/cancel | 查询；取消reason | 执行状态/预算/阶段；202取消请求 |
| POST /runs/{id}/clarifications | R questionSetId/answers[] | 202 runId/inputRevision；过期问题409 |
| GET /cases/{id}；GET /cases/{id}/revisions | O revision；历史列表cursor | 修订及当前ETag；历史按revision倒序 |
| POST /cases/{id}/overlays | R snapshotId/operation/targetRef/reason；O evidenceRefs | 201 overlayRevision；生效需新调查 |
| POST /cases/{id}/submit-review | R revision；O limitationAcknowledgement | 200 reviewStatus；PARTIAL须确认限制 |
| POST /cases/{id}/reviews | R revision/decision/reason | 201 reviewEventId；只审批指定修订 |
| POST /cases/{id}/archive或unarchive | R reason | 200 reviewStatus；解除归档仅Admin |
| POST /cases/{id}/exports | R revision/format | 202 exportId/statusUrl；首版format=DOCX或JSON |
| GET /exports/{id}；GET /exports/{id}/download | 查询；下载 | 状态或鉴权后文件流；没有永久公开地址 |
| GET /runs/{id}/events | O Last-Event-ID请求头 | SSE事件流，受run对应项目权限控制 |
项目成员管理使用PUT /projects/{id}/members/{userId}，请求role和canViewEvidence，响应memberVersion；只允许Admin。source更新使用PUT /sources/{id}带If-Match，不改变已发布快照；source和Case列表分别使用GET /projects/{id}/sources及GET /projects/{id}/cases。GET /snapshots/{id}/manifest提供完整能力及来源摘要。
## 并发和幂等
所有资源更新必须提供If-Match: "v<recordVersion>"；创建资源和运行查询不要求。涉及Case的命令使用CaseRecord版本，澄清及取消使用Run版本；创建新Run同时校验Case ETag。缺少前置条件返回428，过期ETag返回412，活跃Run冲突或业务状态冲突返回409。baseRevision用于指明调查输入，不能替代ETag；不再使用expectedRevision。
所有POST命令要求Idempotency-Key，按actorId、resource、operation隔离，保留24小时；同键同规范化请求返回原状态码和结果，同键异内容409。运行资源本身仍有活跃唯一约束，防止幂等记录过期后重复执行。写事务同时保存幂等结果和资源变更。
分页limit默认50最大200。节点顺序logicalKey/nodeId，历史顺序revision倒序；cursor为服务端签名的不透明值，含scope/filterHash/sortKey，有效期24小时。更改过滤条件或scope后沿用cursor返回400，不能跨项目翻页。
## 错误和事件
错误结构为error.code/message/retryable/details及traceId。400参数错误、401未登录、403无功能权限、404资源不存在或不可见、409业务冲突、412版本过期、422变更语义不支持或快照不可调查、428缺少If-Match、429预算限制、503服务依赖不可用。SQL和密钥不放入details。
SSE事件含eventId/runId/sequence/type/stage/time/payload；type为RUN_STATE、STAGE_CHANGED、TOOL_COMPLETED、BUDGET_UPDATED或RESULT_READY。sequence在Run内单调增长，至少一次投递，客户端按eventId去重；先持久化后推送。事件保留7天，游标过期返回410，客户端重新GET Run后订阅。最终结果以CaseRevision为准，断线不取消Run。
导出状态为QUEUED/RUNNING/SUCCEEDED/FAILED/EXPIRED；默认下载有效期15分钟，每次下载重新鉴权并绑定请求身份，权限撤销立即拒绝。首版通过API代理流式下载，避免无法即时撤销的独立预签名URL。SSE权限撤销后终止连接，缓存和导出均按权限重新检查。
''',True)

put(13,r'''
PostgreSQL保存权威事实与事务，Neo4j保存可重建图投影，满足REQ01、REQ05、REQ06。单机任务队列使用PostgreSQL任务表，不引入额外消息中间件。所有写入通过服务端校验与固定scope；模型不能直接连接两库。
## 数据字典与索引
| 表组 | 主键及关键字段 | 外键和索引 |
| project / project_member / source | UUID主键；member含role/evidencePermission；source含kind/configRevision/credentialRef | member唯一(projectId,userId)；source索引(projectId,id) |
| scan / scan_batch / task_lease | scanId；sourceRefs/configHash/state；batchKey/attemptEpoch/contentHash；leaseOwner/leaseUntil/epoch | batch唯一(scanId,batchKey)；队列索引(state,availableAt)；租约epoch原子递增 |
| snapshot / snapshot_source | snapshotId/projectId/state/quality/manifestHash/sealedAt/publishedAt；各sourceId/versionHash | 唯一(projectId,snapshotId)；来源FK；索引(projectId,publishedAt) |
| arch_node / arch_edge | nodeId及第6章属性；relationId/fromId/toId/kind/condition/confidence/certainty | 复合PK(projectId,snapshotId,id)；端点复合FK；入边与出边索引(scope,toId或fromId,kind) |
| evidence / edge_evidence | evidenceId/sourceHash/snippetHash/location/validation；relationId/evidenceId | scope复合FK；来源位置索引；禁止孤立确定边 |
| unresolved_reference / diagnostic | id/sourceId/location/candidateRefs/code；同scope | 索引(scope,code)；不进入确定图投影 |
| overlay / overlay_event | overlayId/revision/snapshotId/operation/targetRef/actorId/reason | scope FK；唯一(projectId,snapshotId,revision) |
| case_record / case_revision | caseId/recordVersion/latestRevision/reviewStatus；revision和不可变结果JSONB | revision PK(caseId,revision)；snapshot及overlay引用；owner/project索引 |
| investigation_run / run_draft | runId/baseRevision/state/stage/checkpoint/budgetUsed/attemptEpoch；draftVersion/content | 活跃状态caseId部分唯一索引；Run和草案FK |
| tool_call / model_call / run_event | callId/runId/requestHash/response/usage；sequence/type/payload | 唯一(runId,toolCallId)和(runId,sequence)；日志脱敏 |
| outbox / projection_state | eventId/snapshotId/manifestHash/epoch/watermark/state | outbox状态索引；每snapshot唯一投影状态 |
| review_event / export_job / audit_event / idempotency | revision/decision；format/state/expiresAt；actor/action；key/requestHash/response | CaseRevision FK；幂等唯一(actor,resource,operation,key) |
时间为timestamptz，hash为char(64)，业务ID为uuid，状态为受CHECK约束text，契约对象为JSONB并在服务层按schema校验；confidence为numeric(3,2)且0到1。不能只依赖JSONB而取消scope外键。删除project默认软删除；被Case引用的快照禁止普通级联删除。事实表封存后通过数据库角色和服务写入条件拒绝更新，合规删除走独立审计流程。
## Worker租约与批次写入
领取任务在短事务中SELECT FOR UPDATE SKIP LOCKED，设置owner、leaseUntil并递增attemptEpoch。默认租约60秒，每20秒续期；写入使用数据库时间同时检查owner/epoch/leaseUntil及取消标志，过期或已取消返回STALE_ATTEMPT或CANCELLED。检查与写入必须在同一事务持有任务行锁，不能先检查再异步写。
同epoch下AnalysisBatch校验后落到暂存区；允许新epoch复用相同内容hash且已完成的批次。失败批次异内容不覆盖旧记录，记录新attempt并在最终归并时选定有效版本。同键异内容为冲突，需要重建该批次，不让两个Worker混写。故障重试最多2次，间隔1秒和4秒；确定语法错误无重试，模型请求使用第10章独立策略。
## 封存与发布
1  每个source生成固定文件清单；文件变动产生SOURCE_CHANGED并重启该来源扫描，禁止读工作树最新版本补齐。
2  全部预定批次达到终态后，验证必需来源可用、端点绑定、确定边证据及同scope约束。可接受诊断保留并计算quality=PARTIAL；关键来源整体不可用则失败，不发布。
3  在封存事务中锁scan和snapshot，校验当前epoch、全部批次水位及摘要，把有效事实归并完成，snapshot改为SEALED，并写manifest及outbox。大型事实可先分批写入BUILDING区，封存事务只封闭集合和写水位，不要求一次事务搬运全量源码。
4  Projector只能读取SEALED集合，按(scope,id)幂等投影。先节点后边，记录watermark；内容摘要包括排序后的ID、属性、条件和证据引用，不能只比较数量。投影每次写入校验投影epoch；首版每snapshot只允许一个投影任务。
5  数量和摘要一致后，在事务中确认manifestHash未变，projection_state改READY，snapshot改PUBLISHED，再切换activeSnapshotId。封存后分析Worker禁止新增事实；投影失败重试，未发布视图对调查不可见。
取消与发布竞争时锁同一scan/snapshot记录：取消先提交则禁止发布；发布先提交则取消返回409 ALREADY_PUBLISHED。快照state与quality独立，Scan终态SUCCEEDED/PARTIAL映射到PUBLISHED的COMPLETE/PARTIAL。Run固定已发布snapshot和overlay修订，不查询latest替代。
## 投影与恢复
Neo4j节点使用统一ArchNode标签和projectId/snapshotId/nodeId唯一约束；关系唯一键由relationId及scope在投影器中验证，不能假设所有部署版本支持相同关系约束。构建入库校验和投影摘要用于发现重复或缺边。受限邻接模板纳入契约测试，权限条件不可由模型省略。
图库不可用时提供历史报告或503。重建先写独立投影视图，校验后切换，不覆盖正在被查询的已验证投影。PostgreSQL每日备份，试点RPO24小时、RTO4小时；恢复演练必须重放Outbox并核对Case引用和证据，而不是只证明两库能启动。保留被Case引用的源码和证据；合规删除后历史报告标不可复验，不伪造重算。
''',True)

put(14,r'''
## 身份接入和权限实施
服务端采用OIDC身份接入，校验issuer/audience/signature/expiry后映射subject到本地userId；具体企业issuer、clientId和组映射是部署参数。开发可用独立测试身份服务及固定测试账号，不能把前端传来的actorId当身份。服务端project_member为项目授权事实源；模型与Worker使用窄范围服务身份，不继承管理员全域权限。
| 角色 | 允许的操作 | 限制 |
| ProjectAdmin | 来源、成员、扫描、调查及归档管理 | 不因管理员身份默认获得证据查看权；无自动报告审批 |
| Analyst | 授权来源扫描、调查、澄清、标注、送审、草案导出 | 不能批准自己的调查；不能管理成员与密钥 |
| Reviewer | 查看送审修订、批准或否决、导出 | 审批绑定不可变修订；默认禁止自审 |
| Viewer | 查看与导出已批准修订 | 不看草案、Run日志、未授权源码 |
canViewEvidence单独控制片段、search_code结果、模型证据包和包含源码的导出；工具执行权限与证据权限取交集。没有源码权限时模型只能用授权图结构，不得经模型解释旁路泄露。原始证据哈希在隔离区计算，出域前对片段、配置和日志做密钥及个人信息过滤，脱敏后单独计算snippetHash并标明字符位置映射。
部署模型必须确认允许的endpoint与数据范围；默认禁止外发完整源码。API、日志、SSE、导出及模型上下文共用权限过滤，避免单独保护get_source_evidence而在其他通道泄漏。保留策略默认源码随被引用快照保留，普通Run日志30天、事件7天；企业合规策略可以收紧，删除审计和证据不可用标志同步更新。
''')
put(15,r'''
## 五类变更的交付边界
五类变更都提供结构化入口和对象级影响计算；精确修改判断限定第7章能力矩阵。API_CONTRACT_CHANGE首版覆盖字段删除、改名、类型与required变化以及静态路径或方法改变；METHOD_SIGNATURE_CHANGE覆盖已解析直接调用的参数删除、增加和类型变化。不承诺复杂泛型推断、全量数据流或运行时依赖完备。
首个切片必须同时通过当前旧字段引用、拟议别名条件性兼容、已验证映射兼容和动态表名四种样例。没有字段级消费证据时输出UNKNOWN，不能为了完成Demo输出硬编码NO。其余变更类型沿同一契约扩展规则，全部通过第20章后才标MVP完成。
''')

put(16,r'''
实施采用需求到设计到工作包到验收的追踪方式。建议2名开发和1名兼任架构测试人员；原6周仅保留为估算，M0完成样例与兼容性试验后重新排期。以下任务是文档化交付计划，所有状态均为待实施，未执行的验证不能填通过。
## 开发工作包
| 任务与需求 | 交付物及设计位置 | 前置依赖与完成证据 |
| TASK01 REQ01到REQ07 | 样例仓库、来源配置、依赖和版本锁定清单；第7、15章 | 无；列出所有运行与被分析版本，能解析最小样例并记录诊断；先完成M0 |
| TASK02 REQ01/02/03 | 第6、9章契约schema与数据迁移设计 | TASK01；合法/非法样例校验通过，统一字段及scope约束 |
| TASK03 REQ01/05 | PG队列、租约、批次、封存与投影；第4、13章 | TASK02；旧epoch写入被拒，发布前不可见，重放摘要一致 |
| TASK04 REQ01/02 | Java/Spring/MyBatis/PG解析及证据；第7、8章 | TASK02；PAR01到PAR05黄金样例节点/边/位置通过；可与TASK03并行 |
| TASK05 REQ02/03 | Vue请求及字段消费解析；第7章 | TASK02；PAR06/PAR07及动态反例通过；可与TASK04并行 |
| TASK06 REQ03 | 五类传播、兼容合并及评分；第9章 | TASK03/04/05；BR规则、循环、多路径和区间复算通过 |
| TASK07 REQ05/06 | Case/Run/API、权限和UI三工作区；第11、12、14章 | TASK02；可先用契约样例，联调需TASK03/06；版本冲突、澄清、取消和下载撤权通过 |
| TASK08 REQ04 | 模型网关、4阶段模板、工具循环及降级；第10章 | TASK06/07；无效引用、重复工具、拒答、超时、无模型图报告通过 |
| TASK09 REQ05/06 | 导出、审计、保留和恢复；第12到14章 | TASK07；固定修订导出一致，权限撤销生效，恢复引用可定位 |
| TASK10 REQ07 | 独立黄金集、性能/恢复记录、Demo和操作手册；第17、20章 | TASK03到09；所有AC满足，残余限制公开 |
## 阶段与可开工条件
M0完成TASK01和契约评审；M1交付TASK02到04的数据库列到API；M2交付TASK05和完整图发布；M3交付TASK06/07的纯图调查闭环；M4交付TASK08/09；M5执行TASK10。并行仅适用于契约固定后的独立实现，不能跳过集成验证。
技术基线分为产品运行环境和被分析项目矩阵。产品建议Java LTS运行时、Spring Boot MVC、PostgreSQL、Neo4j、Node LTS及Vue3；确切版本、解析器坐标、容器digest、许可证和模型工具能力由TASK01兼容试验锁定，禁止使用浮动latest。选择版本前不建立虚假的“已兼容”清单。源码语言级别与ArchLens自身JDK分别配置。
开工前必须具备可授权的后端与Vue提交、PG schema快照、classpath制品、路由与Profile配置；缺模型配置不阻止确定性内核开发，缺解析依赖则不能承诺符号召回。企业身份issuer、模型endpoint、源码保留策略和业务关键性由实际责任人填写，不由文档代填。
## 部署与交付约定
开发使用单机容器编排，API暴露一个入口，数据库、图库和Worker端口仅内部网络可达；源码卷只读挂载给分析Worker，任务结果卷可写。Worker限CPU、内存、文件数与单文件大小，默认文件数50000、单文件5MB；超限列诊断。API不执行来源代码。生产是否拆服务在试点容量实测后决定。
日志统一traceId/scanId/runId/snapshotId/stage，指标增加租约重领、STALE_ATTEMPT、封存耗时、投影失败、模型重试和UNKNOWN比例。启动先迁移PG并验证Neo4j约束，再启动Worker和API；升级前备份，失败回滚应用与兼容schema迁移，不能删除用户Case作为回滚手段。MVP不承诺在线破坏性schema回滚。
文档交接时可按本章拆成requirements.md、design.md、tasks.md和check_list.md；本DOCX是设计依据，衍生文件不得改变能力边界。需求追踪：REQ01对应TASK01/02/03/04和AC03/07/08；REQ02对应TASK02/04/05及AC01/02/04；REQ03对应TASK02/06及AC05/06；REQ04对应TASK08及AC09/10；REQ05对应TASK03/07/09及AC07/12/13；REQ06对应TASK07/09及AC07/11/14；REQ07对应TASK10及全部AC。
''',True)

put(17,r'''
## 条件性方案的演示方式
演示先锁定before快照并显示拟议别名方案，此时页面结论为UNKNOWN，附“若映射与JSON契约验证通过则预计无需修改”。第二个已验证样例使用独立且明确关联的after快照及验证记录，显示已通过的兼容条件；不得在同一事实图中偷偷替换SQL。现场展示一个页面通过兼容接口A和不兼容接口B同时消费字段，说明节点仍为YES。
''')
put(19,r'''
## 修订后仍需验证的边界
本设计补齐协议和规则，不等于解析器已达到验收指标。最大实施风险仍是跨层字段映射与真实工程依赖解析；应优先在授权真实样例上检验，而不是先开发复杂图形界面。运行版本、制品可用性、企业登录、模型数据范围和业务关键性仍由部署环境决定，TASK01及相关接入任务必须留下确认记录。
''')

put(20,r'''
验收分为设计契约、确定性分析、Agent行为、业务操作和运维恢复。所有检查初始状态为待实施。建立至少40个黄金案例：30个常规案例覆盖5类变更且每类至少6个，另10个动态、歧义或对抗案例；至少一个完整项目作为未参与调参的留出集。两名评审独立标注并裁决，保存标签版本，不使用待测解析器生成标准答案。
## 验收项与证据
| 编号与需求 | 验证内容 | 通过条件及任务 |
| AC01 REQ02 | 节点及直接边 | 支持范围precision≥95%、recall≥90%；分类型报告；TASK04/05 |
| AC02 REQ02 | 跨层路径 | 有效目标路径召回≥90%，缺失逐例解释；TASK04/05/06 |
| AC03 REQ01 | 证据定位 | 确定边100%有同快照有效证据、哈希和规则版本；TASK02/04 |
| AC04 REQ02 | 未知与候选 | 10个反例均不伪报确定，诊断位置和原因完整；TASK04/05 |
| AC05 REQ03 | 修改判断与风险 | 五类规则、拟议兼容、冲突路径、风险区间及HALF_UP复算一致；TASK06 |
| AC06 REQ03 | 资源边界 | 循环终止；超预算返回截断；展示3条不影响已访问判断；TASK06 |
| AC07 REQ01/05/06 | scope和修订 | 越权拒绝；新扫描、overlay、重调查不改旧修订；TASK03/07 |
| AC08 REQ01 | 恢复和发布 | 旧epoch写入被拒；封存后事实不可变；重投影摘要一致；TASK03 |
| AC09 REQ04 | Agent可靠性 | 无效引用强制拦截；语义争议保留；超时或无模型保留图结果；TASK08 |
| AC10 REQ04 | 计划与测试 | 每条建议有理由引用；任务DAG无环；不存在的测试文件不得输出；TASK08 |
| AC11 REQ06 | 只读及注入 | 恶意注释/XML不突破网络、执行和写入白名单；TASK07/09 |
| AC12 REQ05 | 界面闭环 | 消歧、澄清、取消、历史、送审、列表视图、导出均可完成；TASK07/09 |
| AC13 REQ05 | 幂等和事件 | 重复命令不重复Run；过期ETag失败；SSE重连去重、过期恢复；TASK07 |
| AC14 REQ06 | 权限和保留 | 无证据权不经模型/日志/导出泄漏；撤权下载拒绝；删除后标不可复验；TASK07/09 |
## 黄金案例格式和匹配规则
每个fixture包含caseId、sourceManifest、inputFiles、changeSpec、expectedNodes、expectedEdges、expectedPaths、expectedImpacts、expectedDiagnostics及capabilities。期望节点以logicalKey/type匹配，边以from/to/kind/condition/bindingKey匹配，证据同时验证原始位置、哈希及producerVersion；不依赖每次扫描的随机snapshotId。所有预期字段进入版本化JSON，禁止只比较报告自然语言。
影响集合分别比较dependencyPresent、changeRequired、conditionedOutcome和unknown code。路径按标准关系序列或人工定义的等价路径组匹配，不能把任意到达目标的路径当正确。precision=TP/(TP+FP)，recall=TP/(TP+FN)，分母0为N/A；另报告全库成功、不支持、失败文件数和候选边数，不能用缩小支持声明掩盖漏报。
| 示例ID | 输入情形 | 必须得到的结果 |
| FIX01 | 旧SQL引用event_id，提议纯改名且没有兼容措施 | Mapper为YES；上游缺字段证据为UNKNOWN，不能一律YES |
| FIX02 | 只提出global_id AS event_id，未实施验证 | 下游UNKNOWN，conditionedOutcome=NO附条件，兼容state=PROPOSED |
| FIX03 | 有关联验证快照，别名/类型/值语义/JSON均保持，消费覆盖完整 | 对限定消费点NO，原依赖仍可查看 |
| FIX04 | 同页接口A兼容，接口B直接读取被删除字段 | 页面YES，A的NO不覆盖B |
| FIX05 | 页面只读name，接口删eventId；对该页消费点覆盖完整 | 该页NO；增加对象spread未知出口后改为UNKNOWN |
| FIX06 | 动态表名、多个数据源或未知Profile | DYNAMIC_IDENTIFIER或AMBIGUOUS_BINDING，不选任意表 |
| FIX07 | 同路径不同服务或不同produces | 保留不同API节点；请求不足以区分时为候选 |
| FIX08 | 列类型text到uuid，无历史映射验证 | 不宣称等价或回填安全，D未知则风险区间 |
| FIX09 | 图含环和4条不同语义路径，仅展示3条 | 环终止；第4条已访问不兼容路径仍令结论YES |
| FIX10 | 模型引用不存在ID，或存在但内容无关的ID | 不存在ID强制拒；无关语义进入争议复核，不能假设校验器必定识别 |
| FIX11 | WorkerA超时后B接管，A随后写入 | A返回STALE_ATTEMPT；B有效结果与封存摘要不被覆盖 |
| FIX12 | 取消和发布并发，或封存后尝试写事实 | 按提交顺序得到CANCELLED或ALREADY_PUBLISHED；拒绝封存后写入 |
| FIX13 | 同幂等键重复命令、异内容、过期ETag | 同内容同结果；异内容409；过期412；仅一个活跃Run |
| FIX14 | 模型超时或拒答，已有确认ChangeSpec | 图结果保留且标PARTIAL，模型生成段落明确缺失 |
| FIX15 | 元数据采集中并发DDL导致对象查询失败 | 回滚整批重采，不混合尝试；超限失败且无完整快照 |
| FIX16 | 已生成导出后撤销权限，或删除证据 | 下载拒绝；已授权历史报告标证据不可复验 |
FIX表规定必含情形，不代替40个实际fixture；其中正常与对抗分组由标签清单固定，不能重复计数凑数量。每项检查证据记录包含checkId、AC、TASK、版本、执行方式、预期、实际、日志位置、审核人和PASS/FAIL/BLOCKED；当前这些实现证据均未产生。
## 重算与回放
算法重算固定snapshot、overlay、ruleVersion、ChangeSpec、确定的节点/深度/工作单元预算，使用相同排序，比较去掉时间戳的规范化结果哈希。验收时设置足够墙钟时间，单独测试TIME_LIMIT路径；线上5秒超时可能在不同边界截断，不承诺截断集合逐次相同。
历史报告回放直接读取CaseRevision、固定工具结果和已保存模型输出，不再次请求模型，内容哈希应相同。重新调查创建新Run，模型即使配置相同也不要求文字一致；必须重新验证结构、引用和事实约束。原始证据删除后只可回放被允许保留的结果，不能声称可重新计算。
## 性能与验收包
建议8核CPU、16GB内存、SSD、单扫描并发、10万行源码和10万条边以内测量。冷扫描目标≤10分钟；预热节点搜索p95≤1秒、有界Blast p95≤5秒、含模型活跃调查目标≤120秒。至少20次固定查询报告p50/p95/最大值和截断率，p95使用排序后ceil(0.95n)位置；不能通过全部超时返回来满足延迟指标。模型等待、纯算法和数据库耗时分开记录。
验收包包含fixture标签、契约校验、测试结果JSON、引用抽检、UI操作记录、性能及预算日志、故障注入、备份恢复记录、权限反例、版本与许可证清单、操作手册和Demo记录。AC01到AC14全部满足且真实接入参数完成配置后才标MVP完成；文档排版检查不替代这些实施验收。
''',True)

# Append the new authoritative references to the retained appendix.
render('''
## 本次修订核对依据
[R9] PostgreSQL Transaction Isolation\nhttps://www.postgresql.org/docs/current/transaction-iso.html
[R10] MyBatis Mapper XML Files\nhttps://mybatis.org/mybatis-3/sqlmap-xml.html
R2、R9和R10于2026年9月12日复核，用于确认映射条件、事务快照和结果映射机制。其余引用保留原查阅日期。规则权重、超时、范围和预算为本项目设计值，实际兼容版本由TASK01锁定。
## 审查意见落实索引
| 审查项 | 落实位置 | 对应验收 |
| 事实与假设分离 | 第2、9、17章 | AC05及FIX02/03 |
| 字段映射边界 | 第6、7、15章 | AC01/04/05及FIX05 |
| 传播和风险来源 | 第9章 | AC05/06及FIX04/08/09 |
| IR及身份契约 | 第6、13章 | AC03/07及非法契约样例 |
| 解析支持矩阵 | 第7章 | AC01/02/04 |
| 模型调用与复核 | 第10章 | AC09/10及FIX10/14 |
| 租约发布与状态 | 第11、13章 | AC08/12/13及FIX11/12 |
| 数据库一致快照 | 第7、13章 | AC08及FIX15 |
| API及权限闭环 | 第12、14章 | AC07/11/12/13/14 |
| 工作包和可重复验收 | 第16、20章 | TASK01到TASK10及AC01到AC14 |
''')

# Harmonize retained wording with the new state and naming contracts.
appendix.paragraph_format.page_break_before=False
for pp in doc.paragraphs:
    if pp.text.startswith('ADR 01'):
        pp.text='ADR 01：采用Java模块化单体API加独立Java分析Worker和Node前端分析Worker；模块接口为后续拆服务保留边界，执行来源扫描不会占用HTTP请求进程。'
    if pp.text.startswith('ADR 03'): pp.text='ADR 03：单一Investigator状态机调用模型与专业工具。Dependency、Risk、Test、Reviewer是职责划分，具体模型阶段为INTAKE、INVESTIGATE、PLAN、REVIEW，详见第10章。'
    if pp.text.startswith('Collector 不直接写 Neo4j'):
        pp.text='Collector不直接写Neo4j，Analyzer不调用LLM生成确定关系，Graph不承担审批，Case不保存连接密钥。各批次按第6章AnalysisBatch契约交付，Node和Edge端点校验在封存前完成。'
    if pp.text.startswith('控制面接受项目注册'):
        pp.text='控制面接受项目注册、扫描与调查请求；数据面在隔离Worker读取固定来源生成IR，再合并符号、验证证据和写入事实。Agent只通过Tool Gateway读取PUBLISHED快照，通过Model Gateway调用批准模型，源码内容不是控制指令。'
for tt in doc.tables:
    for rr in tt.rows:
        for cc in rr.cells:
            for pp in cc.paragraphs:
                if pp.text=='Case 固定 repository SHA 与 schema 摘要':pp.text='Case固定各来源提交、配置与schema摘要'
                if pp.text=='sourceUri / sourceHash':pp.text='sourceId / sourceHash'
                if pp.text=='kind / producer':pp.text='kind / producerVersion'

# Consistent table formatting, black headings, retained page geometry.
for section in doc.sections:
    section.bottom_margin=Inches(0.9)
    section.footer_distance=Inches(0.35)
for st in doc.styles:
    if st.type==1:
        rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');st.element.rPr.insert(0,rf)
        rf.set(qn('w:eastAsia'),'等线')
        if st.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
            st.font.color.rgb=RGBColor(0,0,0)
for pp in doc.paragraphs:
    if pp.style.name.startswith('Heading'):
        pp.paragraph_format.keep_with_next=True
        for r in pp.runs:r.font.color.rgb=RGBColor(0,0,0)
    pp.paragraph_format.widow_control=True
for ti,tt in enumerate(doc.tables):
    if ti<2:continue
    pr=tt._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    for i,rr in enumerate(tt.rows):
        rp=rr._tr.get_or_add_trPr()
        if i==0 and rp.find(qn('w:tblHeader')) is None:rp.append(OxmlElement('w:tblHeader'))
        for cc in rr.cells:
            cc.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cc._tc.get_or_add_tcPr()
            for old in list(cp):
                if old.tag in [qn('w:shd'),qn('w:tcMar')]:cp.remove(old)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'0B2D4D' if i==0 else ('EFF4F8' if i%2 else 'FFFFFF'));cp.append(sh)
            ma=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');ma.append(e)
            cp.append(ma)
            for pp in cc.paragraphs:
                for r in pp.runs:r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
doc.core_properties.title='ArchLens 企业系统架构取证智能体设计文档'
doc.core_properties.subject='1.1 实施设计修订 基于原v1.0'
out=ROOT/'ArchLens_设计文档_v1.0.docx';doc.save(out)
txt='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
(ROOT/'qa/revised_text.txt').write_text(txt,encoding='utf-8')
assert len([p for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'^\d+  ',p.text)])==20
assert 'UX Case' not in txt
for expected in ['conditionedOutcome','STALE_ATTEMPT','REPEATABLE READ READ ONLY','FRONTEND_ACCESS','ModelGateway','AC14','TASK10','FIX16']:
    assert expected in txt,expected
print(json.dumps({'output':str(out),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'characters':len(txt),'sha256':hashlib.sha256(out.read_bytes()).hexdigest()},ensure_ascii=False))
